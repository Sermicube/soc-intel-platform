"""
app/reports/word_generator.py
─────────────────────────────────────────────────────────────
Generador de reportes en formato Word (.docx) usando python-docx.

Word es útil para reportes que los analistas quieren editar
antes de enviar al cliente.
"""

from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from app.config import get_settings
from app.models.schemas import IOCEnrichmentResult, RiskLevel
from app.utils.logger import logger

RISK_RGB = {
    "critical": RGBColor(0xFF, 0x44, 0x44),
    "high": RGBColor(0xFF, 0x88, 0x00),
    "medium": RGBColor(0xFF, 0xCC, 0x00),
    "low": RGBColor(0x44, 0xBB, 0x44),
    "info": RGBColor(0x44, 0x88, 0xFF),
}


class WordReportGenerator:
    """Genera reportes de incidentes en formato .docx."""

    def __init__(self):
        self.settings = get_settings()

    async def generate(
        self,
        investigation_id: str,
        title: str,
        analyst_name: str,
        ioc_results: list[IOCEnrichmentResult],
        overall_risk_level: str,
        overall_risk_score: int,
        notes: str = "",
        client_name: str | None = None,
    ) -> str:
        """Genera el .docx y retorna la ruta del archivo."""
        doc = Document()

        # ─── Configuración de página ───────────────────────────
        section = doc.sections[0]
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

        self._add_header(doc, title, analyst_name, overall_risk_level, overall_risk_score, client_name)
        self._add_executive_summary(doc, ioc_results, overall_risk_level, overall_risk_score)
        self._add_ioc_table(doc, ioc_results)
        self._add_technical_details(doc, ioc_results)
        self._add_risk_assessment(doc, overall_risk_level)
        self._add_recommendations(doc, overall_risk_level)

        if notes:
            doc.add_heading("6. Notas del Analista", level=1)
            doc.add_paragraph(notes)

        output_dir = self.settings.reports_path
        filename = f"incident_report_{investigation_id[:8]}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        filepath = output_dir / filename

        doc.save(str(filepath))
        logger.info(f"Word generado: {filepath}")
        return str(filepath)

    def _add_header(self, doc, title, analyst_name, risk_level, risk_score, client_name):
        # Clasificación
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("🔒 CONFIDENCIAL — USO INTERNO SOC")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

        # Título principal
        heading = doc.add_heading("REPORTE DE INCIDENTE DE SEGURIDAD", level=0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        subtitle = doc.add_paragraph(title)
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.runs[0].font.size = Pt(14)
        subtitle.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)

        # Metadata
        doc.add_paragraph()
        table = doc.add_table(rows=2, cols=4)
        table.style = "Table Grid"

        cells = [
            ("Fecha:", datetime.now().strftime("%d/%m/%Y %H:%M")),
            ("Analista:", analyst_name),
            ("Nivel de Riesgo:", risk_level.upper()),
            ("Score:", f"{risk_score}/100"),
        ]

        risk_color = RISK_RGB.get(risk_level.lower(), RISK_RGB["info"])
        for i, (label, value) in enumerate(cells):
            row = i // 2
            col = (i % 2) * 2
            table.rows[row].cells[col].text = label
            table.rows[row].cells[col].paragraphs[0].runs[0].font.bold = True
            table.rows[row].cells[col + 1].text = value

        if client_name:
            doc.add_paragraph(f"Cliente: {client_name}")

        doc.add_paragraph()

    def _add_executive_summary(self, doc, ioc_results, risk_level, risk_score):
        doc.add_heading("1. Resumen Ejecutivo", level=1)
        total = len(ioc_results)
        malicious = sum(1 for r in ioc_results if r.risk_level in (RiskLevel.CRITICAL, RiskLevel.HIGH))
        suspicious = sum(1 for r in ioc_results if r.risk_level == RiskLevel.MEDIUM)
        clean = total - malicious - suspicious

        summary = (
            f"Se analizaron {total} indicadores de compromiso (IOCs) utilizando múltiples "
            f"fuentes de inteligencia de amenazas. El análisis reveló {malicious} IOC(s) "
            f"maliciosos o de alto riesgo, {suspicious} sospechosos, y {clean} sin amenazas. "
            f"El nivel de riesgo global es {risk_level.upper()} (Score: {risk_score}/100)."
        )
        doc.add_paragraph(summary)

    def _add_ioc_table(self, doc, ioc_results):
        doc.add_heading("2. Indicadores de Compromiso", level=1)

        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"

        headers = ["IOC", "Tipo", "Riesgo", "Score", "VT Detecciones"]
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            hdr_cells[i].paragraphs[0].runs[0].font.bold = True

        for r in ioc_results:
            row_cells = table.add_row().cells
            vt_det = "N/A"
            if r.virustotal and not r.virustotal.error:
                vt_det = r.virustotal.detection_ratio

            row_cells[0].text = r.ioc_value[:50]
            row_cells[1].text = r.ioc_type.value.upper()
            row_cells[2].text = r.risk_level.value.upper()
            row_cells[3].text = f"{r.risk_score}/100"
            row_cells[4].text = vt_det

        doc.add_paragraph()

    def _add_technical_details(self, doc, ioc_results):
        doc.add_heading("3. Descripción Técnica", level=1)
        for i, r in enumerate(ioc_results, 1):
            doc.add_heading(f"IOC #{i}: {r.ioc_value}", level=2)
            if r.summary:
                doc.add_paragraph(r.summary)
            if r.virustotal and not r.virustotal.error:
                vt = r.virustotal
                doc.add_paragraph(
                    f"VirusTotal: {vt.detection_ratio} motores detectaron amenaza. "
                    f"Reputación: {vt.reputation}."
                )
            if r.abuseipdb and not r.abuseipdb.error:
                ab = r.abuseipdb
                doc.add_paragraph(
                    f"AbuseIPDB: Confianza de abuso {ab.abuse_confidence_score}%. "
                    f"Reportes: {ab.total_reports}. País: {ab.country_code}."
                )
            if r.shodan and not r.shodan.error and r.shodan.ip_str:
                sh = r.shodan
                doc.add_paragraph(
                    f"Shodan: {len(sh.open_ports)} puertos abiertos. "
                    f"Organización: {sh.org or 'Desconocida'}."
                )

    def _add_risk_assessment(self, doc, risk_level):
        doc.add_heading("4. Evaluación de Riesgo", level=1)
        assessments = {
            "critical": "El riesgo es CRÍTICO. Se requiere respuesta inmediata y contención urgente.",
            "high": "El riesgo es ALTO. Se requiere atención urgente e investigación detallada.",
            "medium": "El riesgo es MEDIO. Se requiere monitorización activa y mayor investigación.",
            "low": "El riesgo es BAJO. Se recomienda monitorización estándar.",
            "info": "Sin amenazas identificadas. Se recomienda mantener vigilancia estándar.",
        }
        doc.add_paragraph(assessments.get(risk_level.lower(), assessments["info"]))

    def _add_recommendations(self, doc, risk_level):
        doc.add_heading("5. Acciones Recomendadas", level=1)
        recommendations = {
            "critical": [
                "INMEDIATO: Aislar los sistemas afectados",
                "INMEDIATO: Bloquear IOCs en todos los controles de seguridad",
                "Iniciar proceso formal de respuesta a incidentes",
                "Preservar evidencia forense",
                "Notificar a dirección y equipos legales",
            ],
            "high": [
                "URGENTE: Bloquear IOCs en firewall y EDR",
                "Escalar al equipo de IR",
                "Analizar logs para detectar actividad relacionada",
            ],
            "medium": [
                "Bloquear o monitorizar los IOCs",
                "Añadir IOCs al SIEM para correlación",
                "Investigar origen y contexto",
            ],
            "low": ["Monitorizar IOCs pasivamente", "Registrar para seguimiento"],
            "info": ["Registrar IOCs para referencia futura"],
        }

        for rec in recommendations.get(risk_level.lower(), recommendations["info"]):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(rec)
